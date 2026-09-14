# -*- coding: utf-8 -*-
"""问题2 结果写出：日汇总 CSV、JSON，并填写附件5 的 ``result2.xlsx``（2–12 月）。"""
from __future__ import annotations

import json
import shutil

import pandas as pd
from openpyxl import load_workbook

from common.constants import BLOCK_LABELS, KEY_IDX, KEY_PERIODS, N, PAPER_DATES, XLSX_START
from common.paths import find_xlsx, tables_dir
from common.storage import merge_emergency
from common.time_index import interval_labels


def export_result2(recs):
    labels = interval_labels()
    tpl = find_xlsx("result2")
    dst = tables_dir("q2") / "result2.xlsx"
    shutil.copy(tpl, dst)
    wb = load_workbook(dst)
    ws0, ws1, ws2 = wb.worksheets[0], wb.worksheets[1], wb.worksheets[2]
    bydate = {r["date"]: r for r in recs}
    for i in range(2, ws0.max_row + 1):
        val = ws0.cell(i, 1).value
        if val is None or str(val).strip() in {"⁝", "..."}:
            continue
        dt = pd.Timestamp(val).strftime("%Y-%m-%d")
        if dt not in bydate:
            continue
        r = bydate[dt]
        for t in range(N):
            ws0.cell(i, 2 + t, round(float(r["q"][t]), 4))
        ws0.cell(i, 2 + N, round(float(r["q"].sum()), 4))
        ws0.cell(i, 3 + N, round(float(r["total_cost"]), 4))
    headers1 = [c.value for c in ws1[1]]
    wb.remove(ws1)
    ws1 = wb.create_sheet("充放电量", 1)
    ws1.append(headers1)
    for r in (x for x in recs if x["date"] >= XLSX_START):
        C, D, E = r["C"], r["D"], r["E"]
        for k, lab in enumerate(BLOCK_LABELS):
            sl = slice(k * 24, (k + 1) * 24)
            row = [r["date"] if k == 0 else None, lab, round(float(C[sl].sum()), 4), round(float(D[sl].sum()), 4),
                   "0:00" if k == 0 else ("24:00" if k == 1 else None),
                   round(float(E[0] if k == 0 else E[-1] if k == 1 else 0), 4) if k <= 1 else None]
            if k > 1:
                row[4], row[5] = None, None
            ws1.append(row)
    headers2 = [c.value for c in ws2[1]]
    wb.remove(ws2)
    ws2 = wb.create_sheet("紧急购电量", 2)
    ws2.append(headers2)
    for r in (x for x in recs if x["date"] >= XLSX_START):
        segs = merge_emergency(r["h"], labels)
        if not segs:
            ws2.append([r["date"], "无", 0.0])
        else:
            for i, (lab, qty) in enumerate(segs):
                ws2.append([r["date"] if i == 0 else None, lab, round(qty, 4)])
    wb.save(dst)
    rows = [{"date": r["date"], "Qplan": float(r["q"].sum()), "H": float(r["h"].sum()),
             "plan_cost": r["plan_cost"], "em_cost": r["em_cost"], "total_cost": r["total_cost"],
             "E0": float(r["E"][0]), "E24": float(r["E"][-1])} for r in recs]
    daily = pd.DataFrame(rows)
    daily.to_csv(tables_dir("q2") / "daily_summary.csv", index=False, encoding="utf-8-sig")
    paper = []
    for r in recs:
        if r["date"] not in PAPER_DATES:
            continue
        item = {"日期": r["date"]}
        for name, idx in zip(KEY_PERIODS, KEY_IDX):
            item[name] = round(float(r["q"][idx]), 4)
        item["全天计划购电量"] = round(float(r["q"].sum()), 4)
        item["全天购电费"] = round(float(r["total_cost"]), 4)
        paper.append(item)
    pd.DataFrame(paper).to_csv(tables_dir("q2") / "table_paper_plan.csv", index=False, encoding="utf-8-sig")
    year = {"total_cost": float(daily["total_cost"].sum()), "plan_cost": float(daily["plan_cost"].sum()),
            "em_cost": float(daily["em_cost"].sum())}
    (tables_dir("q2") / "year_summary.json").write_text(json.dumps(year, ensure_ascii=False, indent=2), encoding="utf-8")
    return dst, daily, year
