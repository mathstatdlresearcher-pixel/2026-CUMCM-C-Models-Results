# -*- coding: utf-8 -*-
"""问题1 结果写出：CSV、JSON，并按附件5 模板填写 ``result1.xlsx``。"""
from __future__ import annotations

import json
import shutil

import pandas as pd
from openpyxl import load_workbook

from common.constants import BLOCK_LABELS, KEY_PERIODS, N
from common.paths import find_xlsx, tables_dir


def export_tables(df: pd.DataFrame, sol: dict, report: dict):
    out_dir = tables_dir("q1")
    G, C, D, Suse, E = sol["G"], sol["C"], sol["D"], sol["Suse"], sol["E"]
    out = df.copy()
    out["购电量_kWh"] = G
    out["充电量_kWh"] = C
    out["放电量_kWh"] = D
    out["光伏利用_kWh"] = Suse
    out["弃光_kWh"] = df["pv_kwh"].to_numpy() - Suse
    out["时段初储电量_kWh"] = E[:-1]
    out["时段末储电量_kWh"] = E[1:]
    out["购电费用_元"] = df["price"].to_numpy() * G
    cols = [
        "interval", "price", "load_kw", "pv_kw", "购电量_kWh", "充电量_kWh", "放电量_kWh",
        "光伏利用_kWh", "弃光_kWh", "时段初储电量_kWh", "时段末储电量_kWh", "购电费用_元",
    ]
    out[cols].to_csv(out_dir / "schedule_144.csv", index=False, encoding="utf-8-sig", float_format="%.6f")
    key_rows = [{"时间段": lab, "最优购电量_kWh": float(G[out.index[out["interval"] == lab][0]])} for lab in KEY_PERIODS]
    key_rows.append({"时间段": "全天购电量", "最优购电量_kWh": report["grid_total"]})
    key_rows.append({"时间段": "全天购电费_元", "最优购电量_kWh": report["obj"]})
    pd.DataFrame(key_rows).to_csv(out_dir / "table1_key_periods.csv", index=False, encoding="utf-8-sig")
    block_c, block_d = [], []
    for k in range(6):
        sl = slice(k * 24, (k + 1) * 24)
        block_c.append(float(C[sl].sum()))
        block_d.append(float(D[sl].sum()))
    tab2 = pd.DataFrame({"时间段": BLOCK_LABELS, "充电量_kWh": block_c, "放电量_kWh": block_d})
    tab2.to_csv(out_dir / "table2_4h_charge_discharge.csv", index=False, encoding="utf-8-sig")
    (out_dir / "summary.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    tpl = find_xlsx("result1")
    dst = out_dir / "result1.xlsx"
    shutil.copy(tpl, dst)
    wb = load_workbook(dst)
    ws1, ws2 = wb.worksheets[0], wb.worksheets[1]
    for i in range(N):
        ws1.cell(row=2 + i, column=2, value=round(float(G[i]), 4))
    for k in range(6):
        ws2.cell(row=2 + k, column=2, value=round(block_c[k], 4))
        ws2.cell(row=2 + k, column=3, value=round(block_d[k], 4))
    ws2.cell(row=2, column=5, value=round(float(E[0]), 4))
    ws2.cell(row=3, column=5, value=round(float(E[-1]), 4))
    wb.save(dst)
    return out, tab2
