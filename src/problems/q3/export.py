# -*- coding: utf-8 -*-
"""问题3 结果写出：日汇总与紧急购电时段，填写附件5 的 ``result3.xlsx``。"""
from __future__ import annotations

import json
import shutil

import pandas as pd
from openpyxl import load_workbook

from common.constants import BLOCK_LABELS, N, PAPER_DATES, XLSX_START
from common.paths import find_xlsx, tables_dir
from common.storage import merge_emergency
from common.time_index import interval_labels


def export_result3(recs):
    labels = interval_labels()
    tpl = find_xlsx("result3")
    dst = tables_dir("q3") / "result3.xlsx"
    shutil.copy(tpl, dst)
    wb = load_workbook(dst)
    ws0, ws_adj, ws1, ws2 = wb.worksheets[0], wb.worksheets[1], wb.worksheets[2], wb.worksheets[3]
    bydate = {r["date"]: r for r in recs}

    def fill(ws, key):
        for i in range(2, ws.max_row + 1):
            val = ws.cell(i, 1).value
            if val is None or str(val).strip() in {"⁝", "..."}:
                continue
            dt = pd.Timestamp(val).strftime("%Y-%m-%d")
            if dt not in bydate:
                continue
            q = bydate[dt][key]
            for t in range(N):
                ws.cell(i, 2 + t, round(float(q[t]), 4))
            ws.cell(i, 2 + N, round(float(q.sum()), 4))
            ws.cell(i, 3 + N, round(float(bydate[dt]["total_cost"]), 4))

    fill(ws0, "G0")
    fill(ws_adj, "G_final")
    headers1 = [c.value for c in ws1[1]]
    wb.remove(ws1)
    ws1 = wb.create_sheet("充放电量", 2)
    ws1.append(headers1)
    for r in (x for x in recs if x["date"] >= XLSX_START):
        for k, lab in enumerate(BLOCK_LABELS):
            sl = slice(k * 24, (k + 1) * 24)
            row = [r["date"] if k == 0 else None, lab, round(float(r["C"][sl].sum()), 4), round(float(r["D"][sl].sum()), 4),
                   "0:00" if k == 0 else ("24:00" if k == 1 else None),
                   round(float(r["E"][0] if k == 0 else r["E"][-1] if k == 1 else 0), 4) if k <= 1 else None]
            if k > 1:
                row[4] = row[5] = None
            ws1.append(row)
    headers2 = [c.value for c in ws2[1]]
    wb.remove(ws2)
    ws2 = wb.create_sheet("紧急购电量", 3)
    ws2.append(headers2)
    for r in (x for x in recs if x["date"] >= XLSX_START):
        segs = merge_emergency(r["H"], labels)
        if not segs:
            ws2.append([r["date"], "无", 0.0])
        else:
            for i, (lab, qty) in enumerate(segs):
                ws2.append([r["date"] if i == 0 else None, lab, round(qty, 4)])
    wb.save(dst)
    daily = pd.DataFrame(
        [{"date": r["date"], "Q0": float(r["G0"].sum()), "Qfinal": float(r["G_final"].sum()),
          "H": float(r["H"].sum()), "plan_cost": r["plan_cost"], "adj_cost": r["adj_cost"],
          "em_cost": r["em_cost"], "total_cost": r["total_cost"], "E0": float(r["E"][0]), "E24": float(r["E"][-1])}
         for r in recs]
    )
    daily.to_csv(tables_dir("q3") / "daily_summary.csv", index=False, encoding="utf-8-sig")
    year = {k: float(daily[k].sum()) for k in ["plan_cost", "adj_cost", "em_cost", "total_cost"]}
    (tables_dir("q3") / "year_summary.json").write_text(json.dumps(year, ensure_ascii=False, indent=2), encoding="utf-8")
    return dst, daily, year
